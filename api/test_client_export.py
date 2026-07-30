"""Выгрузка списка клиентов из админки.

Проверяем, что файл каждого формата действительно собирается, содержит данные
клиентов и открывается своей библиотекой. Отдельно — кириллица: без нужного
шрифта PDF молча превратил бы русские буквы в чёрные квадраты.
"""

import io
import re
import zipfile
from decimal import Decimal

from django.contrib.auth.models import User
from django.test import Client, TestCase
from django.urls import reverse

from .models import ClientProfile, Distributor, Region
from .services.exports import CLIENT_COLUMNS, export_clients


class ClientExportTests(TestCase):
    def setUp(self):
        self.distributor = Distributor.objects.create(
            name="Дистрибьютор МСК", inn="1112223334", phone="1", email="d@e.co"
        )
        self.region = Region.objects.create(
            code="77", name="Москва — Центральный", distributor=self.distributor
        )
        manager = User.objects.create_user(username="manager", password="pw")
        manager.first_name, manager.last_name = "Пётр", "Менеджеров"
        manager.save()

        user = User.objects.create_user(
            username="+79001110000", password="pw", email="client@test.ru"
        )
        self.client_profile = ClientProfile.objects.create(
            user=user, inn="5556667778", company_name="ООО «Ромашка»",
            contact_name="Иван Иванов", phone="+79001110000",
            region=self.region, city="Москва", distributor=self.distributor,
            manager=manager, partner_status="Gold", total_purchases=Decimal("2500000"),
        )

    def _queryset(self):
        return ClientProfile.objects.select_related("user", "region", "distributor", "manager")

    # ── Общее ────────────────────────────────────────────────────────────────

    def test_all_formats_produce_a_file(self):
        for fmt in ("xlsx", "docx", "pdf", "csv"):
            with self.subTest(fmt=fmt):
                response = export_clients(self._queryset(), fmt)
                self.assertGreater(len(response.content), 500)
                self.assertIn("attachment;", response["Content-Disposition"])
                self.assertIn(f".{fmt}", response["Content-Disposition"])

    def test_unknown_format_rejected(self):
        with self.assertRaises(ValueError):
            export_clients(self._queryset(), "rtf")

    def test_empty_list_does_not_crash(self):
        ClientProfile.objects.all().delete()
        for fmt in ("xlsx", "docx", "pdf", "csv"):
            with self.subTest(fmt=fmt):
                self.assertGreater(len(export_clients(self._queryset(), fmt).content), 100)

    # ── Excel ────────────────────────────────────────────────────────────────

    def test_xlsx_has_headers_and_data(self):
        from openpyxl import load_workbook

        response = export_clients(self._queryset(), "xlsx")
        sheet = load_workbook(io.BytesIO(response.content)).active

        headers = [cell.value for cell in sheet[1]]
        self.assertEqual(headers, [title for title, _, _ in CLIENT_COLUMNS])

        row = {headers[i]: cell.value for i, cell in enumerate(sheet[2])}
        self.assertEqual(row["Компания"], "ООО «Ромашка»")
        self.assertEqual(row["ИНН"], "5556667778")
        self.assertEqual(row["Ранг"], "Gold")
        self.assertEqual(row["Оборот, ₽"], "2 500 000")
        self.assertEqual(row["Регион"], "Москва — Центральный")
        self.assertEqual(row["Менеджер"], "Пётр Менеджеров")
        self.assertEqual(row["Email"], "client@test.ru")

    def test_xlsx_header_is_frozen(self):
        from openpyxl import load_workbook

        response = export_clients(self._queryset(), "xlsx")
        sheet = load_workbook(io.BytesIO(response.content)).active
        self.assertEqual(sheet.freeze_panes, "A2")

    # ── Word ─────────────────────────────────────────────────────────────────

    def test_docx_contains_client_row(self):
        response = export_clients(self._queryset(), "docx")
        with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
            xml = archive.read("word/document.xml").decode()

        texts = re.findall(r"<w:t[^>]*>([^<]*)</w:t>", xml)
        self.assertIn("Список клиентов AutoTerra", texts)
        self.assertIn("ООО «Ромашка»", texts)
        self.assertIn("Gold", texts)
        self.assertIn("Иван Иванов", texts)

    def test_docx_opens_as_word_document(self):
        from docx import Document

        response = export_clients(self._queryset(), "docx")
        document = Document(io.BytesIO(response.content))

        self.assertEqual(len(document.tables), 1)
        table = document.tables[0]
        self.assertEqual(len(table.rows), 2)  # шапка + один клиент
        self.assertEqual(len(table.columns), len(CLIENT_COLUMNS))

    # ── PDF ──────────────────────────────────────────────────────────────────

    def test_pdf_embeds_cyrillic_font(self):
        response = export_clients(self._queryset(), "pdf")
        data = response.content

        self.assertTrue(data.startswith(b"%PDF"))
        # Без встроенного DejaVu кириллица превратилась бы в квадраты.
        self.assertIn(b"DejaVuSans", data)
        self.assertIn(b"/FontFile2", data)

    def test_pdf_font_registration_survives_repeat_calls(self):
        # Повторная регистрация шрифта не должна падать.
        export_clients(self._queryset(), "pdf")
        response = export_clients(self._queryset(), "pdf")
        self.assertTrue(response.content.startswith(b"%PDF"))

    # ── CSV ──────────────────────────────────────────────────────────────────

    def test_csv_has_bom_and_semicolons(self):
        response = export_clients(self._queryset(), "csv")
        text = response.content.decode("utf-8")

        # BOM — иначе Excel открывает кириллицу кракозябрами.
        self.assertTrue(text.startswith("﻿"))
        lines = text.lstrip("﻿").splitlines()
        self.assertEqual(lines[0].split(";")[0], "Компания")
        self.assertIn("ООО «Ромашка»", lines[1])


class ClientExportAdminTests(TestCase):
    """Кнопки и действия в админке."""

    def setUp(self):
        self.http = Client()
        self.admin = User.objects.create_superuser(
            username="root", password="pw", email="root@test.ru"
        )
        self.http.force_login(self.admin)

        self.distributor = Distributor.objects.create(
            name="Dist", inn="1112223334", phone="1", email="d@e.co"
        )
        self.region = Region.objects.create(code="77", name="Msk", distributor=self.distributor)
        for index, (company, city) in enumerate(
            [("Ромашка", "Москва"), ("Василёк", "Казань")], start=1
        ):
            user = User.objects.create_user(username=f"+7900111000{index}", password="pw")
            ClientProfile.objects.create(
                user=user, inn=f"555666777{index}", company_name=company,
                contact_name="Иван", phone=f"+7900111000{index}",
                region=self.region, city=city, distributor=self.distributor,
            )

    def test_export_button_downloads_everything(self):
        response = self.http.get(
            reverse("admin:api_clientprofile_export", args=["xlsx"])
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("attachment;", response["Content-Disposition"])

        from openpyxl import load_workbook

        sheet = load_workbook(io.BytesIO(response.content)).active
        self.assertEqual(sheet.max_row, 3)  # шапка + два клиента

    def test_export_button_respects_search(self):
        response = self.http.get(
            reverse("admin:api_clientprofile_export", args=["xlsx"]), {"q": "Василёк"}
        )
        self.assertEqual(response.status_code, 200)

        from openpyxl import load_workbook

        sheet = load_workbook(io.BytesIO(response.content)).active
        self.assertEqual(sheet.max_row, 2)  # шапка + один найденный
        self.assertEqual(sheet.cell(row=2, column=1).value, "Василёк")

    def test_export_button_rejects_unknown_format(self):
        response = self.http.get(
            reverse("admin:api_clientprofile_export", args=["rtf"])
        )
        self.assertEqual(response.status_code, 302)  # редирект обратно в список

    def test_admin_action_exports_selected(self):
        selected = ClientProfile.objects.filter(company_name="Ромашка")
        response = self.http.post(
            reverse("admin:api_clientprofile_changelist"),
            {
                "action": "export_pdf",
                "_selected_action": [str(item.pk) for item in selected],
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")

    def test_changelist_shows_export_buttons(self):
        response = self.http.get(reverse("admin:api_clientprofile_changelist"))
        self.assertEqual(response.status_code, 200)

        html = response.content.decode()
        for label in ("Excel", "Word", "PDF", "CSV"):
            self.assertIn(f">{label}</a>", html)
