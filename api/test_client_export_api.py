"""Выгрузка клиентов из мобильного приложения.

Главное здесь — объём выборки: дистрибьютор не должен выгрузить чужих
клиентов, менеджер региона — клиентов соседнего региона, а клиент не должен
получить доступ к выгрузке вовсе.
"""

import io

from django.contrib.auth.models import User
from django.test import Client, TestCase

from .models import AuthToken, ClientProfile, Distributor, Region


def _companies(response):
    from openpyxl import load_workbook

    sheet = load_workbook(io.BytesIO(response.content)).active
    return sorted(row[0] for row in sheet.iter_rows(min_row=2, values_only=True))


class ClientExportApiTests(TestCase):
    def setUp(self):
        self.http = Client()

        # Два дистрибьютора со своими регионами и клиентами.
        self.msk_distributor = self._distributor("Дист МСК", "1112223334", "dist_msk")
        self.kzn_distributor = self._distributor("Дист КЗН", "9998887776", "dist_kzn")
        self.msk = Region.objects.create(code="77", name="Москва", distributor=self.msk_distributor)
        self.kzn = Region.objects.create(code="16", name="Казань", distributor=self.kzn_distributor)

        self._client_profile("Ромашка МСК", "5556667771", self.msk, self.msk_distributor)
        self._client_profile("Василёк МСК", "5556667772", self.msk, self.msk_distributor)
        self._client_profile("Тюльпан КЗН", "5556667773", self.kzn, self.kzn_distributor)

    def _distributor(self, name, inn, username):
        user = User.objects.create_user(username=username, password="pw")
        user.profile.role = "distributor"
        user.profile.save()
        distributor = Distributor.objects.create(
            user=user, name=name, inn=inn, phone="1", email=f"{username}@e.co"
        )
        distributor.token = AuthToken.objects.create(key=f"token-{username}", user=user)
        return distributor

    def _client_profile(self, company, inn, region, distributor):
        user = User.objects.create_user(username=f"+7900{inn[-6:]}", password="pw")
        return ClientProfile.objects.create(
            user=user, inn=inn, company_name=company, contact_name="Иван",
            phone="+79001110000", region=region, city=region.name, distributor=distributor,
        )

    def _manager(self, username, regions=(), is_global=False):
        user = User.objects.create_user(username=username, password="pw")
        user.profile.role = "admin" if is_global else "manager"
        user.profile.save()
        for region in regions:
            region.manager = user
            region.save(update_fields=["manager"])
        return AuthToken.objects.create(key=f"token-{username}", user=user)

    def _export(self, token, fmt="xlsx", **params):
        query = "&".join(f"{key}={value}" for key, value in {"format": fmt, **params}.items())
        return self.http.get(
            f"/api/clients/export/?{query}",
            HTTP_AUTHORIZATION=f"Bearer {token.key}",
        )

    # ── Доступ ───────────────────────────────────────────────────────────────

    def test_anonymous_rejected(self):
        self.assertEqual(self.http.get("/api/clients/export/").status_code, 401)

    def test_client_cannot_export(self):
        profile = ClientProfile.objects.first()
        token = AuthToken.objects.create(key="client-token", user=profile.user)
        self.assertEqual(self._export(token).status_code, 403)

    # ── Объём выборки ────────────────────────────────────────────────────────

    def test_distributor_gets_only_own_clients(self):
        response = self._export(self.msk_distributor.token)
        self.assertEqual(response.status_code, 200, response.content[:200])
        self.assertEqual(_companies(response), ["Василёк МСК", "Ромашка МСК"])

    def test_other_distributor_sees_only_theirs(self):
        response = self._export(self.kzn_distributor.token)
        self.assertEqual(_companies(response), ["Тюльпан КЗН"])

    def test_region_manager_limited_to_own_regions(self):
        token = self._manager("manager_msk", regions=[self.msk])
        response = self._export(token)
        self.assertEqual(_companies(response), ["Василёк МСК", "Ромашка МСК"])

    def test_global_manager_gets_everyone(self):
        token = self._manager("manager_main", is_global=True)
        response = self._export(token)
        self.assertEqual(
            _companies(response), ["Василёк МСК", "Ромашка МСК", "Тюльпан КЗН"]
        )

    # ── Форматы и фильтры ────────────────────────────────────────────────────

    def test_every_format_downloads(self):
        for fmt, mime in (
            ("xlsx", "spreadsheetml"),
            ("docx", "wordprocessingml"),
            ("pdf", "application/pdf"),
            ("csv", "text/csv"),
        ):
            with self.subTest(fmt=fmt):
                response = self._export(self.msk_distributor.token, fmt=fmt)
                self.assertEqual(response.status_code, 200)
                self.assertIn(mime, response["Content-Type"])
                self.assertIn(f".{fmt}", response["Content-Disposition"])

    def test_unknown_format_is_rejected(self):
        response = self._export(self.msk_distributor.token, fmt="rtf")
        self.assertEqual(response.status_code, 400)
        self.assertIn("Неизвестный формат", response.json()["detail"])

    def test_default_format_is_excel(self):
        response = self.http.get(
            "/api/clients/export/",
            HTTP_AUTHORIZATION=f"Bearer {self.msk_distributor.token.key}",
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])

    def test_search_narrows_the_file(self):
        response = self._export(self.msk_distributor.token, search="Ромашка")
        self.assertEqual(_companies(response), ["Ромашка МСК"])

    def test_search_cannot_reach_other_distributor(self):
        response = self._export(self.msk_distributor.token, search="Тюльпан")
        self.assertEqual(_companies(response), [])

    def test_filter_by_partner_status(self):
        ClientProfile.objects.filter(company_name="Ромашка МСК").update(partner_status="Gold")
        response = self._export(self.msk_distributor.token, partnerStatus="Gold")
        self.assertEqual(_companies(response), ["Ромашка МСК"])
