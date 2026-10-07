from datetime import datetime
from decimal import Decimal
from io import BytesIO
from zoneinfo import ZoneInfo

from django.test import override_settings
from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook

from .models import AuthToken, ClientProfile, Distributor, Payment, YooKassaRegistryImport, YooKassaRegistryOperation
from .test_bonus_account import _Base


@override_settings(SECURE_SSL_REDIRECT=False)
class FinanceTests(_Base):
    url = '/api/finance/yookassa/'
    dates = {'date_from': '2026-10-01', 'date_to': '2026-10-31'}

    def setUp(self):
        super().setUp()
        for role in ('admin', 'distributor', 'manager', 'courier'):
            user = User.objects.create_user(username=role)
            user.profile.role = role
            user.profile.save()
            AuthToken.objects.create(user=user, key=role)
            if role == 'distributor':
                self.distributor.user = user
                self.distributor.save()
        self.foreign = Distributor.objects.create(name='Foreign', inn='8888888888')
        self.order = self._order('1200', 'paid')
        # Actual collected amount, after bonuses/discounts, is the accounting base.
        self.payment = Payment.objects.create(order=self.order, provider_payment_id='pay-1', amount='1000', status='succeeded',
            paid_at=datetime(2026, 10, 5, 12, tzinfo=ZoneInfo('Europe/Moscow')))
        self.other_order = self._order(800, 'paid')
        self.other_order.distributor = self.foreign
        self.other_order.save()
        self.other_payment = Payment.objects.create(order=self.other_order, provider_payment_id='foreign-pay', amount='800', status='succeeded', paid_at=self.payment.paid_at)

    def auth(self, role='admin'):
        return {'HTTP_AUTHORIZATION': f'Bearer {role}'}

    def report(self, role='admin', **kwargs):
        return self.http.get(self.url, {**self.dates, **kwargs}, **self.auth(role))

    def upload(self, text, role='admin', name='registry.csv', **fields):
        raw = text.encode('utf-8') if isinstance(text, str) else text
        return self.http.post(self.url + 'import/', {'file': SimpleUploadedFile(name, raw), **fields}, **self.auth(role))

    def sale(self, pid='pay-1', amount='1000', fee='30', net='970', day='05.10.2026'):
        return f'Идентификатор платежа;Сумма платежа;Сумма комиссии;Сумма за вычетом комиссии;Время платежа\n{pid};{amount};{fee};{net};{day} 12:00:00\n'

    def refund(self, amount='200', rid='refund-1', pid='pay-1', day='10.10.2026'):
        return f'Идентификатор возврата;Идентификатор платежа;Сумма возврата;Валюта возврата;Время возврата\n{rid};{pid};{amount};RUB;{day} 12:00:00\n'

    def test_scope_all_endpoints_and_filter_tampering(self):
        for role in ('manager', 'courier', 'client-token'):
            self.assertEqual(self.report(role).status_code, 403)
            self.assertEqual(self.upload(self.sale(), role).status_code, 403)
            self.assertEqual(self.http.get(self.url + 'export/', self.dates, **self.auth(role)).status_code, 403)
        for path in ('', 'export/', 'import/'):
            method = self.http.post if path == 'import/' else self.http.get
            self.assertEqual(method(self.url + path).status_code, 401)
        self.assertEqual(self.report('distributor', distributor_id=self.foreign.pk).status_code, 403)
        self.assertEqual(self.upload(self.sale(), 'distributor', distributor_id=self.foreign.pk).status_code, 403)
        self.assertEqual(self.http.get(self.url + 'export/', {**self.dates, 'distributor_id': self.foreign.pk}, **self.auth('distributor')).status_code, 403)
        self.assertEqual(self.report('distributor').json()['summary']['sales'], '1000.00')
        self.assertEqual(self.report().json()['summary']['sales'], '1800.00')
        self.assertEqual(self.report(distributor_id=self.foreign.pk).json()['summary']['sales'], '800.00')

    def test_known_fees_refunds_and_repeat_import_do_not_change_payments(self):
        response = self.upload(self.sale(), 'distributor')
        self.assertEqual(response.status_code, 201, response.content)
        self.assertEqual(self.upload(self.sale(), 'distributor').json()['duplicates'], 1)
        self.assertEqual(self.upload(self.refund(), 'distributor').status_code, 201)
        self.assertEqual(self.upload(self.refund(), 'distributor').json()['duplicates'], 1)
        summary = self.report('distributor').json()['summary']
        self.assertEqual(summary['paidOrders'], 1)
        self.assertEqual(summary['sales'], '1000.00')
        self.assertEqual(summary['commission'], '30.00')
        self.assertEqual(summary['refunds'], '200.00')
        self.assertEqual(summary['net'], '770.00')
        self.assertEqual(summary['status'], 'matched')
        self.payment.refresh_from_db()
        self.order.refresh_from_db()
        self.assertEqual(self.payment.amount, Decimal('1000'))
        self.assertEqual(self.payment.status, 'succeeded')
        self.assertEqual(self.order.total_amount, Decimal('1200'))
        self.assertEqual(self.order.status, 'paid')

    def test_unknown_fee_is_not_zero_and_other_providers_excluded(self):
        for provider in ('cash', 'bonus'):
            Payment.objects.create(order=self.order, provider=provider, amount=100, status='succeeded', paid_at=self.payment.paid_at)
        Payment.objects.create(order=self.order, amount=500, status='pending', paid_at=self.payment.paid_at)
        summary = self.report('distributor').json()['summary']
        self.assertEqual(summary['sales'], '1000.00')
        self.assertIsNone(summary['commission'])
        self.assertIsNone(summary['net'])
        self.assertEqual(summary['status'], 'discrepancies')

    def test_amount_unknown_id_and_net_discrepancies(self):
        self.upload(self.sale(amount='999', net='969'), 'distributor')
        self.upload(self.sale(pid='missing'), 'distributor')
        report = self.report('distributor').json()
        self.assertEqual({row['status'] for row in report['discrepancies']}, {'amount_mismatch', 'payment_not_found'})
        self.assertEqual(report['summary']['sales'], '1000.00')
        self.assertIsNone(report['summary']['net'])
        self.assertFalse(self.report(distributor_id=self.foreign.pk).json()['discrepancies'][0]['providerPaymentId'] == 'missing')
        self.assertEqual(self.upload(self.sale(net='960'), 'distributor').status_code, 400)
        YooKassaRegistryOperation.objects.all().delete()
        self.upload(self.sale(net='960'), 'distributor')
        self.assertEqual(self.report('distributor').json()['discrepancies'][0]['status'], 'net_mismatch')

    def test_foreign_import_row_is_skipped_and_never_leaks(self):
        result = self.upload(self.sale(pid='foreign-pay', amount='800', net='770'), 'distributor').json()
        self.assertEqual(result['skipped'], 1)
        self.assertEqual(result['imported'], 0)
        self.assertFalse(YooKassaRegistryOperation.objects.exists())
        self.assertNotIn('foreign-pay', str(self.report('distributor').json()))

    def test_official_preamble_vat_and_cp1251(self):
        text = 'РЕЕСТР ПЛАТЕЖЕЙ ПО ДОГОВОРУ\nДата платежей: 2026-10-05\nИдентификатор платежа;Сумма платежа;Сумма комиссии без НДС;НДС с комиссии;Сумма за вычетом комиссии и НДС;Время платежа\npay-1;1000;30;6,60;963,40;05.10.2026 12:00:00\n\nСумма принятых платежей: 1000 RUB\nПо договору ЭК.1\n'
        self.assertEqual(self.upload(text.encode('cp1251'), 'distributor').status_code, 201)
        self.assertEqual(self.report('distributor').json()['summary']['commission'], '36.60')

    def test_xlsx_payments_and_refunds_all_sheets(self):
        book = Workbook()
        for sheet, text in [(book.active, self.sale()), (book.create_sheet('Возвраты'), self.refund())]:
            for row in text.strip().splitlines():
                sheet.append(row.split(';'))
        output = BytesIO()
        book.save(output)
        response = self.upload(output.getvalue(), 'distributor', 'statement.xlsx')
        self.assertEqual(response.status_code, 201, response.content)
        self.assertEqual(response.json()['imported'], 2)
        self.assertEqual(self.report('distributor').json()['summary']['net'], '770.00')

    def test_period_refund_of_previous_period_and_legacy_dates(self):
        self.payment.paid_at = datetime(2026, 9, 5, 12, tzinfo=ZoneInfo('Europe/Moscow'))
        self.payment.save()
        self.upload(self.refund(), 'distributor')
        summary = self.report('distributor').json()['summary']
        self.assertEqual(summary['sales'], '0.00')
        self.assertEqual(summary['paidOrders'], 0)
        self.assertEqual(summary['net'], '-200.00')
        self.payment.paid_at = None
        self.payment.save()
        Payment.objects.filter(pk=self.payment.pk).update(created_at=datetime(2026, 10, 1, tzinfo=ZoneInfo('Asia/Bishkek')))
        self.assertEqual(self.report('distributor').json()['summary']['sales'], '1000.00')
        self.assertEqual(self.report('distributor', date_to='2026-09-30').status_code, 400)

    def test_multiple_partial_refunds_and_over_refund(self):
        self.upload(self.sale(), 'distributor')
        self.upload(self.refund('200'), 'distributor')
        self.upload(self.refund('300', 'refund-2'), 'distributor')
        self.assertEqual(self.report('distributor').json()['summary']['refunds'], '500.00')
        self.upload(self.refund('600', 'refund-3'), 'distributor')
        report = self.report('distributor').json()
        self.assertIsNone(report['summary']['net'])
        self.assertEqual(report['summary']['refunds'], '0.00')
        self.assertEqual(len([row for row in report['discrepancies'] if row['status'] == 'refund_exceeds_payment']), 3)

    def test_invalid_import_is_atomic_and_conflicting_duplicate_rejected(self):
        bad = self.sale() + 'second;NaN;1;9;05.10.2026 12:00:00\n'
        self.assertEqual(self.upload(bad).status_code, 400)
        self.assertFalse(YooKassaRegistryImport.objects.exists())
        self.upload(self.sale())
        response = self.upload(self.sale(amount='1001', net='971'))
        self.assertEqual(response.status_code, 400)
        self.assertEqual(YooKassaRegistryImport.objects.count(), 1)
        self.assertEqual(YooKassaRegistryOperation.objects.count(), 1)

    def test_order_hint_matching_and_unknown_id_does_not_fallback(self):
        text = f'provider_payment_id;order_id;amount;fee;net;date\n;{self.order.pk};1000;30;970;2026-10-05T12:00:00+03:00\n'
        response = self.upload(text, 'distributor')
        self.assertEqual(response.status_code, 201, response.content)
        self.assertEqual(YooKassaRegistryOperation.objects.get().payment_id, self.payment.pk)
        self.upload(text.replace(f';{self.order.pk};', f'unknown;{self.order.pk};'), 'distributor')
        self.assertTrue(any(row['status'] == 'payment_not_found' for row in self.report('distributor').json()['discrepancies']))

    def test_excel_export_totals_scope_numbers_and_formula_injection(self):
        self.upload(self.sale(), 'distributor')
        self.client_profile.company_name = '=HYPERLINK("bad")'
        ClientProfile.objects.filter(pk=self.client_profile.pk).update(company_name=self.client_profile.company_name)
        response = self.http.get(self.url + 'export/', self.dates, **self.auth('distributor'))
        self.assertEqual(response.status_code, 200)
        book = load_workbook(BytesIO(response.content))
        self.assertEqual(book.sheetnames, ['Итоги', 'Операции', 'Расхождения'])
        self.assertEqual(book['Операции'].max_row, 2, list(book['Операции'].values))
        self.assertEqual(book['Операции']['F2'].value, 1000)
        self.assertEqual(book['Операции']['I2'].value, 970)
        self.assertEqual(book['Операции']['C2'].data_type, 's')
        self.assertEqual(book['Операции'].freeze_panes, 'A2')
        self.assertIsNotNone(book['Операции'].auto_filter.ref)

    def test_order_only_csv_and_overlapping_id_statement_deduplicate(self):
        text = f'order_id,amount,fee,net,date\n{self.order.pk},1000,30,970,2026-10-05T12:00:00+03:00\n'
        response = self.upload(text, 'distributor')
        self.assertEqual(response.status_code, 201, response.content)
        self.assertEqual(self.upload(self.sale(), 'distributor').json()['duplicates'], 1)
        self.assertEqual(YooKassaRegistryOperation.objects.count(), 1)

    def test_missing_payment_is_linked_by_repeat_import(self):
        text = self.sale(pid='late-pay')
        self.upload(text, 'distributor')
        self.assertIsNone(YooKassaRegistryOperation.objects.get().payment_id)
        Payment.objects.create(order=self.order, provider_payment_id='late-pay', status='succeeded', amount=1000, paid_at=self.payment.paid_at)
        self.assertEqual(self.upload(text, 'distributor').json()['duplicates'], 1)
        self.assertIsNotNone(YooKassaRegistryOperation.objects.get().payment_id)

    def test_pagination_full_totals_unique_order_count_and_export(self):
        for index in range(101):
            Payment.objects.create(order=self.order, provider_payment_id=f'page-{index}', status='succeeded', amount=1, paid_at=self.payment.paid_at)
        first = self.report('distributor').json()
        second = self.report('distributor', page=2).json()
        self.assertEqual(first['count'], 102)
        self.assertEqual(len(first['operations']), 100)
        self.assertEqual(len(second['operations']), 2)
        self.assertEqual(first['summary']['paidOrders'], 1)
        self.assertEqual(first['summary']['sales'], '1101.00')
        self.assertEqual(first['summary'], second['summary'])
        response = self.http.get(self.url + 'export/', self.dates, **self.auth('distributor'))
        book = load_workbook(BytesIO(response.content))
        self.assertEqual(book['Операции'].max_row, 103)
        self.assertEqual(book['Расхождения'].max_row, 103)

    def test_invalid_files_and_ambiguous_order_do_not_match(self):
        self.assertEqual(self.upload(b'broken', name='broken.xlsx').status_code, 400)
        self.assertEqual(self.upload(self.sale(), name='file.exe').status_code, 400)
        for amount in ('NaN', '-1', '0', '1.001'):
            self.assertEqual(self.upload(self.sale(amount=amount)).status_code, 400)
        Payment.objects.create(order=self.order, status='succeeded', amount=1000, provider_payment_id='another')
        text = f'order_id;amount;fee;net;date\n{self.order.pk};1000;30;970;2026-10-05T12:00:00+03:00\n'
        self.assertEqual(self.upload(text, 'distributor').status_code, 201)
        self.assertIsNone(YooKassaRegistryOperation.objects.get().payment_id)
        self.assertTrue(any(row['status'] == 'payment_not_found' for row in self.report('distributor').json()['discrepancies']))

    def test_missing_order_payment_repeat_import_uses_original_entry(self):
        order = self._order(700)
        text = f'order_id;amount;fee;net;date\n{order.pk};700;20;680;2026-10-05T12:00:00+03:00\n'
        self.assertEqual(self.upload(text, 'distributor').status_code, 201)
        self.assertIsNone(YooKassaRegistryOperation.objects.get().payment_id)
        payment = Payment.objects.create(order=order, amount=700, status='succeeded', provider_payment_id='late-order-pay', paid_at=self.payment.paid_at)
        result = self.upload(text, 'distributor')
        self.assertEqual(result.status_code, 201, result.content)
        self.assertEqual(result.json()['duplicates'], 1)
        row = YooKassaRegistryOperation.objects.get()
        self.assertEqual(row.payment_id, payment.pk)
        self.assertEqual(row.event_key, 'payment:late-order-pay')
