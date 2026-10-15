"""Regression fixtures reproduce the bands/columns/outline of the real 1C price.

No commercial workbook is stored in Git. Its A/M/N/O/Q/S layout and folder
outline (folders also carry codes and aggregate stock) are reproduced below.
"""
import io
import os
from unittest import skipUnless
from decimal import Decimal

from django.test import TestCase
from openpyxl import Workbook

from .models import MAX_PRODUCT_IMAGES, Distributor, Product
from .services.product_import import parse_products_workbook
from .services.product_categories import CANONICAL_PRODUCT_CATEGORIES, determine_product_category
from .test_product_import import StockUploadFileEndpointTests


def one_c_workbook(rows=None, *, photo_header=True):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Лист_1'
    sheet['A1'] = 'Прайс-лист на 8 октября 2026 г.'
    for col, value in {'A': 'Номенклатура', 'M': 'Артикул', 'N': 'Код', 'Q': '*Основной склад (единый)'}.items():
        sheet[f'{col}3'] = value
    sheet['O4'], sheet['O5'] = 'руб.', 'Включает НДС'
    for col, value in {'O': 'Цена', 'P': 'Валюта', 'Q': 'Остаток', 'R': 'Наличие',
                       'T': 'Источник фото', 'U': 'Статус фото', 'V': 'Кандидат фото'}.items():
        sheet[f'{col}6'] = value
    if photo_header:
        sheet['S6'] = 'Фото'
    if rows is None:
        rows = [
            {'name': '000.Akzo Nobel', 'code': '00023179', 'level': 0},
            {'name': 'Akzo Nobel Оборудование', 'code': '00023484', 'level': 1, 'price': 0, 'quantity': 3},
            {'name': '2002956 Спектрофотометр Automatchic Vision', 'article': '2002956', 'code': '00023078',
             'price': 264471.894, 'quantity': 2, 'status': 'Под заказ', 'images': 'https://example.com/photo.jpg',
             'candidate': 'https://example.com/candidate.jpg', 'level': 2},
        ]
    columns = {'name': 'A', 'article': 'M', 'code': 'N', 'price': 'O', 'quantity': 'Q',
               'status': 'R', 'images': 'S', 'candidate': 'V'}
    for number, row in enumerate(rows, 7):
        sheet.row_dimensions[number].outlineLevel = row.get('level', 0)
        for field, col in columns.items():
            if field in row:
                sheet[f'{col}{number}'] = row[field]
        sheet[f'N{number}'].number_format = row.get('code_format', 'General')
    stream = io.BytesIO()
    workbook.save(stream)
    stream.seek(0)
    stream.name = '1c-price.xlsx'
    return stream


def item(**changes):
    result = {'name': 'Грунт', 'article': 'ARTICLE-1', 'code': '000001', 'price': 100,
              'quantity': 5, 'images': 'https://example.com/a.jpg', 'level': 1}
    result.update(changes)
    return result


class OneCCategoryTests(TestCase):
    def test_requested_product_and_folder_classification(self):
        cases = [
            ('SUNMIGHT Абразивные круги 150мм', '', 'Абразивы'),
            ('Круг для снятия ржавчины', 'Насадки на дрель', 'Абразивы'),
            ('Бокорезы', '', 'Инструменты'),
            ('Плоскогубцы', '', 'Инструменты'),
            ('Монтировки', '', 'Инструменты'),
            ('Гайковерты', '', 'Инструменты'),
            ('Компрессометр', '', 'Инструменты'),
            ('Рефрактометр', '', 'Инструменты'),
            ('Разбавитель', '', 'Растворители'),
            ('Фильтры', '', 'Расходники'),
            ('Лаки для автомобиля', '', 'Лаки'),
            ('Грунтовки для автомобиля', '', 'Грунтовки'),
            ('Шпатлевки для автомобиля', '', 'Шпатлёвки'),
            ('Автомобильная шпатлевка', '', 'Шпатлёвки'),
            ('Краска базовая', '', 'Краски базовые'),
            ('Акриловая эмаль 2К', '', 'Краски 2K'),
            ('Антикор', '', 'Антикоры'),
            ('Герметик', '', 'Герметики'),
            ('Фильтр для краски', 'Базовые эмали', 'Расходники'),
            ('Отвердитель для лака', 'Лаки', 'Отвердители'),
            ('Грунт акриловый 2К', 'Краски 2K', 'Грунтовки'),
            ('Краскопульт для грунта', 'Грунтовки', 'Инструменты'),
            ('Полировальный круг 150мм', '', 'Абразивы'),
            ('Малярный комбинезон XL', 'Разное', 'Расходники'),
            ('Фильтр для лака', 'Лаки', 'Расходники'),
            ('Фильтр для растворителя', '', 'Расходники'),
            ('Шлифовальная машина', '', 'Инструменты'),
            ('Артикул 123', 'SUNMIGHT Абразивные круги 150мм', 'Абразивы'),
            ('Артикул 123', 'Бокорезы', 'Инструменты'),
        ]
        for name, folder, expected in cases:
            with self.subTest(name=name, folder=folder):
                self.assertEqual(determine_product_category(name, folder), expected)

    def test_canonical_categories_are_single_source_and_idempotent(self):
        self.assertEqual(len(CANONICAL_PRODUCT_CATEGORIES), 12)
        self.assertEqual(len(set(CANONICAL_PRODUCT_CATEGORIES)), 12)
        for category in CANONICAL_PRODUCT_CATEGORIES:
            self.assertEqual(determine_product_category(category), category)

    def test_hierarchy_fallback_and_unknown_category(self):
        self.assertEqual(determine_product_category('Материал X', 'Серия X',
            ['Материалы', 'Грунтовки', 'Серия X']), 'Грунтовки')
        self.assertIsNone(determine_product_category('Неизвестный товар', 'Новая папка', ['Разное']))


class OneCParserTests(TestCase):
    def test_real_multilevel_header_name_code_article_price_stock_and_photo(self):
        products, errors = parse_products_workbook(one_c_workbook())
        self.assertEqual(errors, [])
        self.assertEqual(len(products), 1)  # coded folders must not become products
        product = products[0]
        self.assertEqual(product['name'], '2002956 Спектрофотометр Automatchic Vision')
        self.assertEqual(product['sku'], '00023078')
        self.assertEqual(product['wb_article'], '2002956')
        self.assertEqual(product['price'], Decimal('264471.89'))
        self.assertEqual(product['quantity'], 2)
        self.assertEqual(product['images'], ['https://example.com/photo.jpg'])
        self.assertEqual(product['group_name'], '000.Akzo Nobel')
        self.assertEqual(product['category'], 'Инструменты')
        self.assertEqual(product['status'], 'onOrder')
        self.assertTrue(product['_has_price'])
        self.assertTrue(product['_has_quantity'])

    def test_nested_folders_aggregate_stock_and_sibling_categories(self):
        rows = [
            {'name': 'Материалы', 'code': 'G0', 'level': 0},
            {'name': 'Лаки', 'code': 'G1', 'level': 1, 'price': 0, 'quantity': 100},
            {'name': '2K', 'code': 'G2', 'level': 2, 'price': 0, 'quantity': 99},
            item(name='Лак', code='L1', level=3),
            {},  # empty separators do not erase the current folder
            item(name='Лак без артикула', article='', code='L2', level=3),
            {'name': 'Грунты', 'code': 'G3', 'level': 1},
            item(code='P1', level=2),
            {'name': 'Инструменты', 'code': 'G4', 'level': 0},
            item(name='Гайковерт', code='T1', level=1),
        ]
        products, errors = parse_products_workbook(one_c_workbook(rows))
        self.assertEqual(errors, [])
        self.assertEqual([p['sku'] for p in products], ['L1', 'L2', 'P1', 'T1'])
        self.assertEqual([p['category'] for p in products], ['Лаки', 'Лаки', 'Грунтовки', 'Инструменты'])
        self.assertEqual([p['group_name'] for p in products], ['Материалы'] * 3 + ['Инструменты'])
        self.assertEqual(products[1]['wb_article'], '')

    def test_all_workbook_products_have_canonical_categories(self):
        rows = []
        for index, (folder, name) in enumerate([
            ('SUNMIGHT Абразивные круги 150мм', 'Круг P120'),
            ('Насадки на дрель', 'Круг для снятия ржавчины'),
            ('Бокорезы', 'Модель X'), ('Плоскогубцы', 'Модель X'),
            ('Монтировки', 'Модель X'), ('Гайковерты', 'Модель X'),
            ('Растворители', 'Разбавитель'), ('Разное', 'Фильтры'),
            ('Лаки для автомобиля', 'Материал X'),
            ('Грунтовки для автомобиля', 'Материал X'),
            ('Шпатлевки для автомобиля', 'Материал X'),
        ]):
            rows += [{'name': folder, 'code': f'GROUP-{index}', 'level': 0},
                     item(name=name, code=f'PRODUCT-{index}', level=1)]
        products, errors = parse_products_workbook(one_c_workbook(rows))
        self.assertEqual(errors, [])
        self.assertEqual(len(products), 11)
        self.assertTrue(all(product['category'] in CANONICAL_PRODUCT_CATEGORIES for product in products))
        self.assertEqual(products[1]['category'], 'Абразивы')
        self.assertEqual(products[1]['group_name'], 'Насадки на дрель')

    def test_unknown_category_skips_row_with_warning_instead_of_raw_folder(self):
        products, errors = parse_products_workbook(one_c_workbook([
            {'name': 'Неизвестная папка', 'code': 'GROUP', 'level': 0},
            item(name='Неизвестный товар', code='UNKNOWN'),
            item(name='Лак', code='GOOD'),
        ]))
        self.assertEqual([product['sku'] for product in products], ['GOOD'])
        self.assertEqual(products[0]['category'], 'Лаки')
        self.assertEqual(len(errors), 1)
        self.assertIn('Строка 8 (UNKNOWN)', errors[0])
        self.assertIn('не удалось определить категорию', errors[0])
        self.assertIn('Неизвестная папка', errors[0])

    def test_explicit_one_c_category_is_normalized(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(['Номенклатура', 'Код', 'Категория', 'Цена', 'Остаток'])
        sheet.append(['Модель X', 'CODE', 'Бокорезы ' + 'x' * 500, 100, 3])
        stream = io.BytesIO()
        workbook.save(stream)
        stream.seek(0)
        products, errors = parse_products_workbook(stream)
        self.assertEqual(errors, [])
        self.assertEqual(products[0]['category'], 'Инструменты')

    def test_fractional_stock_floors_with_row_warning_and_localized_numbers(self):
        products, errors = parse_products_workbook(one_c_workbook([
            item(price='1\u00a0234,565', quantity='2,656'),
            item(code='000002', price=0.609, quantity=0.971),
        ]))
        self.assertEqual([p['price'] for p in products], [Decimal('1234.57'), Decimal('0.61')])
        self.assertEqual([p['quantity'] for p in products], [2, 0])
        self.assertEqual(len(errors), 2)
        self.assertIn('Строка 7 (000001)', errors[0])
        self.assertIn('округлён вниз до 2', errors[0])
        self.assertIn('округлён вниз до 0', errors[1])

    def test_invalid_rows_skip_without_stopping_valid_rows(self):
        bad = [
            item(code=None), item(code='BAD-NAME', name=None), item(code='BAD-PRICE', price='не число'),
            item(code='NEGATIVE', quantity=-1), item(code='INFINITY', price='Infinity'),
            item(code='NAN', quantity='NaN'), item(code='BIG-PRICE', price='10000000000'),
            item(code='BIG-QTY', quantity=2147483648), item(code='LONG-CODE-' + 'x' * 65),
            item(code='LONG-NAME', name='x' * 256),
            item(code='#REF!'), item(code='ERROR-NAME', name='#DIV/0!'),
        ]
        products, errors = parse_products_workbook(one_c_workbook(bad + [item(code='GOOD')]))
        self.assertEqual([p['sku'] for p in products], ['GOOD'])
        self.assertEqual(len(errors), len(bad))
        self.assertTrue(all('пропущено' in warning for warning in errors))
        self.assertIn('нет кода', errors[0])

    def test_multiple_photos_newlines_duplicates_invalid_urls_and_limit(self):
        links = [f'https://example.com/{i}.jpg' for i in range(20)]
        value = '\n'.join(links) + ';' + links[0] + ',ftp://example.com/not-a-photo'
        products, errors = parse_products_workbook(one_c_workbook([item(images=value)]))
        self.assertEqual(products[0]['images'], links[:MAX_PRODUCT_IMAGES])
        self.assertEqual(len(errors), 2)
        self.assertTrue(any('сохранены первые 15' in warning for warning in errors))
        self.assertTrue(any('без http(s)-ссылки' in warning for warning in errors))

    def test_numeric_code_zero_padding_and_zero_stock(self):
        products, errors = parse_products_workbook(one_c_workbook([
            item(code=123, code_format='00000000', article=987, quantity=0),
            item(code=0, article='000ABC', quantity=1),
        ]))
        self.assertEqual(errors, [])
        self.assertEqual(products[0]['sku'], '00000123')
        self.assertEqual(products[0]['wb_article'], '987')
        self.assertEqual(products[1]['sku'], '0')
        self.assertEqual(products[1]['wb_article'], '000ABC')
        self.assertEqual(products[0]['quantity'], 0)

    def test_missing_photo_column_does_not_clear_images(self):
        products, errors = parse_products_workbook(one_c_workbook(photo_header=False))
        self.assertEqual(errors, [])
        self.assertNotIn('images', products[0])

    def test_flat_one_c_and_column_order_are_detected_by_names(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(['Фото', 'Артикул', 'Остаток', 'Номенклатура', 'Цена', 'Код'])
        sheet.append(['https://example.com/a.jpg', 'SUPPLIER-1', 4, 'Лак', 350.5, '000CODE'])
        stream = io.BytesIO()
        workbook.save(stream)
        stream.seek(0)
        products, errors = parse_products_workbook(stream)
        self.assertEqual(errors, [])
        self.assertEqual(products[0]['sku'], '000CODE')
        self.assertEqual(products[0]['wb_article'], 'SUPPLIER-1')
        self.assertEqual(products[0]['images'], ['https://example.com/a.jpg'])

    def test_one_c_without_code_header_does_not_fall_back_to_article(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(['Номенклатура', 'Артикул', 'Цена', 'Остаток'])
        sheet.append(['Лак', 'SUPPLIER', 100, 2])
        stream = io.BytesIO()
        workbook.save(stream)
        stream.seek(0)
        products, errors = parse_products_workbook(stream)
        self.assertEqual(products, [])
        self.assertIn('«Код»', errors[0])

    def test_ordinary_excel_keeps_article_as_sku(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(['Артикул', 'Код', 'Наименование', 'Цена', 'Остаток', 'Категория'])
        sheet.append(['LEGACY', 'EXTERNAL', 'Лак', 200, 3, 'Произвольная категория'])
        stream = io.BytesIO()
        workbook.save(stream)
        stream.seek(0)
        products, errors = parse_products_workbook(stream)
        self.assertEqual(errors, [])
        self.assertEqual(products[0]['sku'], 'LEGACY')
        self.assertEqual(products[0]['category'], 'Произвольная категория')


class OneCUploadEndpointTests(TestCase):
    setUp = StockUploadFileEndpointTests.setUp
    _upload = StockUploadFileEndpointTests._upload

    def test_repeat_import_updates_same_distributor_code_and_catalog_images(self):
        existing = Product.objects.create(distributor=self.distributor, sku='00023078', name='Старое название',
            category='Старый раздел', brand='Бренд', price=1, quantity=99)
        other = Distributor.objects.create(name='Другой', inn='901234567891')
        foreign = Product.objects.create(distributor=other, sku='00023078', name='Чужой товар',
            category='Категория', price=10, quantity=11)
        for _ in range(2):
            response = self._upload(one_c_workbook())
            self.assertEqual(response.status_code, 200, response.content)
            self.assertEqual(response.json()['created'], 0)
            self.assertEqual(response.json()['updated'], 1)
            self.assertEqual(response.json()['processed'], 1)
            existing.refresh_from_db()
            self.assertEqual(existing.name, '2002956 Спектрофотометр Automatchic Vision')
            self.assertEqual(existing.price, Decimal('264471.89'))
            self.assertEqual(existing.quantity, 2)
            self.assertEqual(existing.wb_article, '2002956')
            self.assertEqual(existing.brand, 'Бренд')
            self.assertEqual(existing.category, 'Инструменты')
            self.assertIn(existing.category, CANONICAL_PRODUCT_CATEGORIES)
        self.assertEqual(Product.objects.filter(distributor=self.distributor).count(), 1)
        foreign.refresh_from_db()
        self.assertEqual(foreign.price, Decimal('10'))
        self.assertEqual(foreign.quantity, 11)
        auth = {'HTTP_AUTHORIZATION': f'Bearer {self.token.key}'}
        response = self.http.get('/api/distributor/stock/', **auth)
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()['results'][0]['images'], ['https://example.com/photo.jpg'])

    def test_reimport_changed_price_stock_article_photos_without_duplicates(self):
        response = self._upload(one_c_workbook())
        self.assertEqual(response.json()['created'], 1)
        product = Product.objects.get(distributor=self.distributor, sku='00023078')
        response = self._upload(one_c_workbook([item(code='00023078', article='NEW-ARTICLE', name='Новый грунт',
            price='123,45', quantity=7, images='https://example.com/new.jpg')]))
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()['created'], 0)
        self.assertEqual(response.json()['updated'], 1)
        product.refresh_from_db()
        self.assertEqual(product.name, 'Новый грунт')
        self.assertEqual(product.price, Decimal('123.45'))
        self.assertEqual(product.quantity, 7)
        self.assertEqual(product.wb_article, 'NEW-ARTICLE')
        self.assertEqual(product.images, ['https://example.com/new.jpg'])
        self.assertEqual(product.category, 'Грунтовки')
        self.assertIn(product.category, CANONICAL_PRODUCT_CATEGORIES)
        self.assertEqual(Product.objects.filter(distributor=self.distributor).count(), 1)

    def test_bad_rows_and_fractional_warnings_are_returned_by_endpoint(self):
        response = self._upload(one_c_workbook([item(code='BAD', price='ошибка'), item(quantity='0,971')]))
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()['created'], 1)
        self.assertEqual(len(response.json()['errors']), 2)
        self.assertFalse(Product.objects.filter(sku='BAD').exists())
        self.assertEqual(Product.objects.get(sku='000001').quantity, 0)

    def test_unknown_category_does_not_create_or_overwrite_existing_product(self):
        original = Product.objects.create(distributor=self.distributor, sku='EXISTING',
            name='Лак', category='Лаки', price=100, quantity=5)
        response = self._upload(one_c_workbook([
            {'name': 'Неизвестная папка', 'code': 'GROUP', 'level': 0},
            item(name='Неизвестный товар', code='NEW'),
            item(name='Неизвестный товар', code='EXISTING', price=500, quantity=20),
            item(name='Грунт', code='VALID'),
        ]))
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()['created'], 1)
        self.assertEqual(response.json()['updated'], 0)
        self.assertEqual(len(response.json()['errors']), 2)
        self.assertFalse(Product.objects.filter(distributor=self.distributor, sku='NEW').exists())
        original.refresh_from_db()
        self.assertEqual(original.category, 'Лаки')
        self.assertEqual(original.name, 'Лак')
        self.assertEqual(original.price, Decimal('100'))
        self.assertEqual(original.quantity, 5)

    def test_only_unknown_categories_returns_existing_upload_error_contract(self):
        response = self._upload(one_c_workbook([item(name='Неизвестный товар', code='UNKNOWN')]))
        self.assertEqual(response.status_code, 400)
        self.assertIn('не удалось определить категорию', response.json()['detail'])
        self.assertEqual(len(response.json()['errors']), 1)
        self.assertFalse(Product.objects.filter(distributor=self.distributor).exists())

    def test_missing_photo_column_preserves_existing_photos_on_update(self):
        original = Product.objects.create(distributor=self.distributor, sku='00023078', name='Товар',
            category='Категория', images=['https://example.com/preserved.jpg'])
        response = self._upload(one_c_workbook(photo_header=False))
        self.assertEqual(response.status_code, 200, response.content)
        original.refresh_from_db()
        self.assertEqual(original.images, ['https://example.com/preserved.jpg'])


class ExpandedOneCCategoryTests(TestCase):
    def test_group_families_and_common_abbreviations(self):
        cases = [
            ('Кисть флейцевая', 'ОРМИС', 'Инструменты'),
            ('Набор шпателей ПВХ', 'ОРМИС', 'Инструменты'),
            ('Нож Стандарт', 'ОРМИС', 'Инструменты'),
            ('Набор щеток д/зачис.поверх.пластик', 'ОРМИС', 'Абразивы'),
            ('INP MBX щетка большая', 'INP', 'Абразивы'),
            ('Щетка-крацовка д/УШМ', 'ОРМИС', 'Абразивы'),
            ('Щиток токаря,защитный лицевой', 'ОРМИС', 'Расходники'),
            ('Круг лепестковый P80', 'DEBEVER', 'Абразивы'),
            ('Шлиф.блок 120х98х13мм P120', 'SUNMIGHT', 'Абразивы'),
            ('Ручной шлифовальный блок-ракель без абразива', 'SANDWOX', 'Инструменты'),
            ('Круг полир.150х25мм', 'РМ Разное', 'Абразивы'),
            ('Водостойкая бумагаP1000', 'SIA', 'Абразивы'),
            ('Перф.листы Goldflex', 'MIRKA Перфорированные листы в рулоне Goldflex Soft', 'Абразивы'),
            ('Полоса IRIDIUM 81*133мм', 'MIRKA Полосы IRIDIUM 81х133мм 54 отв', 'Абразивы'),
            ('Блок подготовки воздуха', 'WIEDER KRAFT', 'Инструменты'),
            ('Цилиндр WDK-серия', 'WIEDER KRAFT', 'Инструменты'),
            ('Набор для рихтовки', 'WIEDER KRAFT', 'Инструменты'),
            ('Stirring Lid-Assembied', 'PPG SELEMIX', 'Инструменты'),
            ('Бесщеточная орбитальная машинка', 'TesKom Машинки', 'Инструменты'),
            ('Ручка системы X', 'РМ Столы/стойки', 'Инструменты'),
            ('Комплект панелей для верстака', 'FERRUM (Металлическая мебель)', 'Инструменты'),
            ('Модель изоляции X', 'Шумоизоляция К-УРАЛ', 'Расходники'),
            ('Валик прикаточный', 'Шумоизоляция К-УРАЛ', 'Инструменты'),
            ('Полумаска противогазоаэрозольная', 'JP Полумаска', 'Расходники'),
            ('Предфильтр', '3M Китай', 'Расходники'),
            ('Фреза для ремонта пластика', 'FULLEN-POLYMER', 'Инструменты'),
            ('Синий плоский профиль', 'FULLEN-POLYMER', 'Расходники'),
            ('Набор треугольных прутов сварочных для PP1', 'BAMPERUS', 'Расходники'),
            ('Холодная сварка универсальная', 'ОРМИС', 'Герметики'),
            ('Уплотнительная масса под кисть', 'POLFILL', 'Герметики'),
            ('Антисиликон NF', 'Fitter Разное', 'Растворители'),
            ('LE-series SRA 1л', 'LENERGY', 'Растворители'),
            ('Смывка краски (бутылка)', 'Смывка', 'Растворители'),
            ('Преобразователь ржавчины с распылителем', 'ПОЛИХИМ', 'Антикоры'),
            ('Праймер (с кисточкой)', '3M Китай', 'Грунтовки'),
            ('Базовая краска с тест-картой', 'Краски', 'Краски базовые'),
            ('Отвердитель для лака', 'Лаки', 'Отвердители'),
            ('Грунт с отвердителем', 'Грунты и отвердители', 'Грунтовки'),
            ('Аэрозольная акриловая эмаль 1К', 'NC Спрей-краска', 'Краски базовые'),
            ('Эмаль для суппортов RAPTOR, аэрозоль', 'U-POL Термостойкие краски HIGH HEAT!', 'Краски базовые'),
            ('Защитное покрытие RAPTOR', 'U-POL Raptor', 'Краски 2K'),
            ('Эмаль MobiHARD ПУ', 'ЭМАЛИ', 'Краски 2K'),
            ('Эмаль HARDMAX PUR', 'ЭМАЛИ', 'Краски 2K'),
            ('SCV BT LV 151 DTM КОНВ.П/М', 'Sikkens Тонеры Autocoat BT SCV', 'Краски 2K'),
            ('Пигмент Hymax RED LUMINA', 'Hymax Спецпродукты', 'Краски базовые'),
            ('Паста колер. POLYMER U', 'Пигменты PALIZH', 'Краски базовые'),
            ('SYSTEMIX Пигментная паста', 'ПИГМЕНТНЫЕ ПАСТЫ SYSTEMIX', 'Краски базовые'),
            ('Пигментная паста 2К', 'Пигменты PALIZH', 'Краски 2K'),
            ('Сухое проявочное покрытие', 'Разное', 'Расходники'),
            ('Силикагель для осушителя', 'РМ Разное', 'Расходники'),
            ('Паста полировальная', 'HIQ', 'Расходники'),
            ('Жидкость для детейлинга', 'MENZERNA', 'Расходники'),
        ]
        for name, folder, expected in cases:
            with self.subTest(name=name, folder=folder):
                self.assertEqual(determine_product_category(name, folder), expected)

    def test_unknown_products_in_mixed_groups_remain_unresolved(self):
        for folder in ['Разное', '001.Краски', 'ОРМИС', 'WIEDER KRAFT', 'FULLEN-POLYMER', 'LENERGY']:
            with self.subTest(folder=folder):
                self.assertIsNone(determine_product_category('Неизвестный товар XYZ', folder))
        self.assertIsNone(determine_product_category('Цилиндр XYZ', 'Неизвестная папка'))

    def test_audit_preserves_full_hierarchy_and_does_not_write_database(self):
        import json
        import tempfile
        from pathlib import Path
        from django.core.management import call_command
        rows = [
            {'name': 'Материалы', 'code': 'G0', 'level': 0},
            {'name': 'Шумоизоляция', 'code': 'G1', 'level': 1},
            {'name': 'Шумоизоляция К-УРАЛ', 'code': 'G2', 'level': 2},
            item(name='Модель X', code='P1', level=3),
            item(name='Модель Y', code='P2', level=3),
        ]
        audit = []
        products, errors = parse_products_workbook(one_c_workbook(rows), category_audit=audit)
        self.assertEqual(errors, [])
        self.assertEqual(len(products), 2)
        self.assertEqual(audit[0]['hierarchy'], ['Материалы', 'Шумоизоляция', 'Шумоизоляция К-УРАЛ'])
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            workbook = root / 'price.xlsx'
            workbook.write_bytes(one_c_workbook(rows).getvalue())
            baseline = root / 'baseline.json'
            baseline.write_text(json.dumps([dict(row, category=None) for row in audit]), encoding='utf-8')
            call_command('audit_1c_categories', str(workbook), output_dir=root / 'report',
                baseline=baseline, expect_total=2, require_complete=True, stdout=io.StringIO())
            report = json.loads((root / 'report/category-audit.json').read_text(encoding='utf-8'))
            self.assertEqual(report['summary']['classified'], 2)
            self.assertEqual(report['summary']['unresolved'], 0)
            self.assertEqual(report['summary']['unresolved_before'], 2)
            self.assertEqual(report['unresolved_before'][0]['folder_total'], 2)
            self.assertEqual(report['unresolved_before'][0]['folder_unresolved_before'], 2)
        self.assertFalse(Product.objects.exists())

    def test_unknown_rows_are_audited_without_default_category(self):
        audit = []
        products, errors = parse_products_workbook(
            one_c_workbook([item(name='Неизвестный товар', code='UNKNOWN')]), category_audit=audit)
        self.assertEqual(products, [])
        self.assertEqual(len(errors), 1)
        self.assertEqual(len(audit), 1)
        self.assertIsNone(audit[0]['category'])

    def test_expanded_groups_repeat_import_without_duplicates(self):
        from .services.product_import import upsert_products
        distributor = Distributor.objects.create(name='Category audit', inn='901234567892')
        rows = [item(name='Шпатель фасадный', code='TOOL'),
                item(name='Предфильтр для полумаски', code='SUPPLY'),
                item(name='SRA для переходов', code='SOLVENT')]
        products, errors = parse_products_workbook(one_c_workbook(rows))
        self.assertEqual(errors, [])
        self.assertEqual(upsert_products(distributor, products), (3, 0))
        changed, errors = parse_products_workbook(one_c_workbook([
            dict(row, quantity=9, price=200) for row in rows]))
        self.assertEqual(errors, [])
        self.assertEqual(upsert_products(distributor, changed), (0, 3))
        self.assertEqual(Product.objects.filter(distributor=distributor).count(), 3)
        self.assertEqual(set(Product.objects.filter(distributor=distributor).values_list('category', flat=True)),
                         {'Инструменты', 'Расходники', 'Растворители'})
        self.assertTrue(all(p.quantity == 9 and p.price == Decimal('200')
                            for p in Product.objects.filter(distributor=distributor)))


@skipUnless(os.environ.get('AUTOTERRA_TEST_1C_WORKBOOK'), 'Private real workbook path is not configured')
class RealOneCWorkbookTests(TestCase):
    def test_complete_real_workbook_and_two_upserts(self):
        from pathlib import Path
        from .services.product_import import upsert_products
        audit = []
        products, errors = parse_products_workbook(
            Path(os.environ['AUTOTERRA_TEST_1C_WORKBOOK']), category_audit=audit)
        self.assertEqual(len(audit), 3968)
        self.assertEqual(len(products), 3968)
        self.assertFalse(any('не удалось определить категорию' in error for error in errors))
        self.assertTrue(all(product['category'] in CANONICAL_PRODUCT_CATEGORIES for product in products))
        self.assertEqual(len({product['sku'] for product in products}), 3968)
        distributor = Distributor.objects.create(name='Private real workbook test', inn='901234567893')
        self.assertEqual(upsert_products(distributor, products), (3968, 0))
        self.assertEqual(upsert_products(distributor, products), (0, 3968))
        self.assertEqual(Product.objects.filter(distributor=distributor).count(), 3968)
        for product in Product.objects.filter(distributor=distributor):
            self.assertIn(product.category, CANONICAL_PRODUCT_CATEGORIES)
