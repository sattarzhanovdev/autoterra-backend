"""Internal YooKassa reconciliation. Provider statements are evidence, not payment commands."""
import csv
import hashlib
import io
import json
import re
from datetime import datetime
from decimal import Decimal, InvalidOperation
from zoneinfo import ZoneInfo
from zipfile import BadZipFile, ZipFile
from xml.etree.ElementTree import ParseError

from django.db import transaction
from django.db.models import Q
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from ..models import Payment, YooKassaRegistryImport, YooKassaRegistryOperation

STATUS_LABELS = {
    'matched': 'Сверено', 'empty': 'Нет операций', 'discrepancies': 'Есть расхождения',
    'not_in_registry': 'Нет в реестре', 'payment_not_found': 'Платеж не найден',
    'amount_mismatch': 'Сумма не совпадает', 'order_mismatch': 'Заказ не совпадает',
    'currency_mismatch': 'Валюта не совпадает', 'payment_not_succeeded': 'Оплата не подтверждена',
    'fee_unknown': 'Комиссия неизвестна', 'net_mismatch': 'Сумма к зачислению не совпадает',
    'refund_exceeds_payment': 'Возвраты превышают оплату',
}
ZERO = Decimal('0.00')
MAX_ROWS = 20000
ALIASES = {
    'payment_id': ['provider_payment_id', 'payment_id', 'идентификатор платежа', 'номер транзакции'],
    'refund_id': ['refund_id', 'идентификатор возврата'],
    'order_id': ['order_id', 'номер заказа', 'id заказа'],
    'amount': ['amount', 'сумма платежа'],
    'refund': ['refund', 'refund_amount', 'сумма возврата'],
    'fee': ['fee', 'commission', 'комиссия', 'сумма комиссии', 'сумма комиссии без ндс'],
    'vat': ['fee_vat', 'ндс с комиссии'],
    'receipt_fee': ['комиссия за чеки (с ндс)'],
    'net': ['net', 'net_amount', 'сумма к зачислению', 'сумма за вычетом комиссии', 'сумма за вычетом комиссии и ндс'],
    'date': ['date', 'occurred_at', 'дата', 'время платежа', 'время возврата'],
    'currency': ['currency', 'валюта платежа', 'валюта возврата'],
}


def norm(value):
    return re.sub(r'\s+', ' ', str(value or '').replace('\ufeff', '').replace('\xa0', ' ')).strip().lower()


def money(value):
    try:
        result = Decimal(str(value).replace('\xa0', '').replace(' ', '').replace(',', '.'))
        if not result.is_finite() or result < 0 or result >= Decimal('10000000000') or result != result.quantize(Decimal('.01')):
            raise ValueError()
        return result.quantize(Decimal('.01'))
    except (InvalidOperation, ValueError):
        raise ValueError('Некорректная денежная сумма') from None


def moment(value):
    if isinstance(value, datetime):
        result = value
    else:
        raw = str(value).strip()
        try:
            result = datetime.fromisoformat(raw.replace('Z', '+00:00'))
        except ValueError:
            result = None
            for fmt in ('%d.%m.%Y %H:%M:%S', '%d.%m.%Y %H:%M', '%d.%m.%Y'):
                try:
                    result = datetime.strptime(raw, fmt)
                    break
                except ValueError:
                    pass
            if result is None:
                raise ValueError('Некорректная дата операции')
    # YooKassa statements use Moscow time; ISO timestamps may specify their zone.
    return result if timezone.is_aware(result) else result.replace(tzinfo=ZoneInfo('Europe/Moscow'))


def tables(content, filename):
    if filename.lower().endswith('.xlsx'):
        try:
            with ZipFile(io.BytesIO(content)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 50 * 1024 * 1024:
                    raise ValueError('Слишком большой распакованный XLSX')
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except (BadZipFile, OSError, KeyError, ValueError, ParseError, InvalidFileException) as exc:
            raise ValueError('Поврежденный XLSX') from exc
        try:
            for sheet in book:
                yield sheet.iter_rows(values_only=True)
        finally:
            book.close()
    elif filename.lower().endswith('.csv'):
        try:
            raw = content.decode('utf-8-sig')
        except UnicodeDecodeError:
            raw = content.decode('cp1251')
        # Header may follow statement title/period, which have no delimiters.
        header = next((line for line in raw.splitlines() if any(a in norm(line) for a in ALIASES['payment_id'] + ALIASES['order_id'])), '')
        delimiter = max((';', ',', '\t'), key=header.count)
        yield csv.reader(io.StringIO(raw), delimiter=delimiter)
    else:
        raise ValueError('Поддерживаются CSV и XLSX')


def parse_statement(content, filename):
    result = []
    visited = 0
    for table in tables(content, filename):
        columns = None
        for number, cells in enumerate(table, 1):
            visited += 1
            if visited > MAX_ROWS:
                raise ValueError(f'Не более {MAX_ROWS} строк в файле')
            names = [norm(v) for v in cells]
            header = {key: next((i for i, name in enumerate(names) if name in aliases), None) for key, aliases in ALIASES.items()}
            if (header['payment_id'] is not None or header['order_id'] is not None) and (header['amount'] is not None or header['refund'] is not None):
                if header['amount'] is not None and header['refund'] is not None:
                    raise ValueError('Платежи и возвраты должны быть отдельными таблицами или листами')
                columns = header
                continue
            if columns is None or not any(v is not None and str(v).strip() for v in cells):
                continue
            # Official statements include free-text footers, outside the operation table.
            if len(cells) == 1 or norm(cells[0]).startswith(('сумма ', 'число ', 'по договору', 'total ', 'number of ')):
                continue
            def get(key):
                index = columns[key]
                return cells[index] if index is not None and index < len(cells) else None
            try:
                refund = columns['refund'] is not None and columns['amount'] is None
                kind = 'refund' if refund else 'payment'
                pid = str(get('payment_id') or '').strip()
                rid = str(get('refund_id') or '').strip()
                hint_value = str(get('order_id') or '').strip()
                if hint_value and not re.fullmatch(r'[0-9]+(?:\.0+)?', hint_value):
                    raise ValueError('Некорректный order_id')
                hint = int(Decimal(hint_value)) if hint_value else None
                if hint is not None and not 0 < hint <= 9223372036854775807:
                    raise ValueError('Некорректный order_id')
                if not pid and not hint:
                    raise ValueError('Нужен ID платежа или order_id')
                if len(pid) > 128 or len(rid) > 128:
                    raise ValueError('Слишком длинный идентификатор')
                if refund and not rid:
                    raise ValueError('Для возврата нужен уникальный refund_id')
                amount = money(get('refund') if refund else get('amount'))
                if amount <= 0:
                    raise ValueError('Сумма операции должна быть больше нуля')
                date = moment(get('date'))
                fee = money(get('fee')) if get('fee') not in (None, '') else None
                if fee is not None:
                    fee += money(get('vat') or 0) + money(get('receipt_fee') or 0)
                net = money(get('net')) if get('net') not in (None, '') else None
                if not refund:
                    if fee is None and net is not None:
                        fee = amount - net
                    if net is None and fee is not None:
                        net = amount - fee
                    if (fee is not None and fee < 0) or (net is not None and net < 0):
                        raise ValueError('Комиссия превышает оплату')
                currency = str(get('currency') or 'RUB').strip().upper()
                if currency != 'RUB':
                    raise ValueError('Учет поддерживает только RUB')
                row = dict(kind=kind, provider_payment_id=pid, provider_refund_id=rid, order_id_hint=hint,
                           occurred_at=date, amount=amount, fee=fee, net=net, currency=currency)
                row['event_key'] = f'{kind}:{rid if refund else pid or "order-" + str(hint)}'
                row['fingerprint'] = hashlib.sha256(json.dumps(row, default=str, sort_keys=True).encode()).hexdigest()
                result.append(row)
            except (ValueError, TypeError, OverflowError) as exc:
                raise ValueError(f'Строка {number}: {exc}') from exc
    if not result:
        raise ValueError('В файле не найдены операции с поддерживаемыми заголовками')
    return result


def match_payment(row):
    payments = Payment.objects.filter(provider='yookassa').select_related('order__distributor', 'order__client')
    if row.provider_payment_id:
        payments = payments.filter(provider_payment_id=row.provider_payment_id)
    elif row.order_id_hint:
        payments = payments.filter(order_id=row.order_id_hint, status='succeeded')
    else:
        return None
    candidates = list(payments[:2])
    return candidates[0] if len(candidates) == 1 else None


def fingerprint(row, payment=None):
    canonical = {key: getattr(row, key) for key in ('kind', 'provider_payment_id', 'provider_refund_id',
        'order_id_hint', 'amount', 'fee', 'net', 'currency')}
    if payment:
        if row.order_id_hint == payment.order_id:
            canonical['order_id_hint'] = None
        if not row.provider_payment_id:
            canonical['provider_payment_id'] = payment.provider_payment_id
    canonical['occurred_at'] = row.occurred_at.astimezone(ZoneInfo('UTC')).isoformat()
    return hashlib.sha256(json.dumps(canonical, default=str, sort_keys=True).encode()).hexdigest()


@transaction.atomic
def import_statement(content, filename, user, distributor):
    rows = parse_statement(content, filename)
    batch = YooKassaRegistryImport.objects.create(uploaded_by=user, distributor=distributor,
        filename=filename[:255], file_hash=hashlib.sha256(content).hexdigest())
    skipped = 0
    for data in rows:
        row = YooKassaRegistryOperation(**data)
        payment = match_payment(row)
        if distributor and payment and payment.order.distributor_id != distributor.pk:
            skipped += 1
            continue
        owner = payment.order.distributor if payment else distributor
        original_key = row.event_key
        if payment and not row.provider_payment_id and payment.provider_payment_id:
            row.provider_payment_id = payment.provider_payment_id
            if row.kind == 'payment':
                row.event_key = f'payment:{payment.provider_payment_id}'
        row.fingerprint = fingerprint(row, payment)
        existing = YooKassaRegistryOperation.objects.select_for_update().filter(event_key__in=[row.event_key, original_key]).first()
        if existing:
            if distributor and existing.distributor_id != distributor.pk:
                skipped += 1
                continue
            if fingerprint(existing, payment) != row.fingerprint:
                raise ValueError('Повторная операция содержит другие данные. Импорт отменен; проверьте реестр.')
            # Retry resolves a row whose payment did not exist during the first import.
            if payment and existing.payment_id is None:
                existing.payment = payment
                existing.distributor = owner
                existing.event_key = row.event_key
                existing.provider_payment_id = row.provider_payment_id
                existing.fingerprint = row.fingerprint
                existing.save(update_fields=['payment', 'distributor', 'event_key', 'provider_payment_id', 'fingerprint'])
            batch.duplicate_count += 1
            continue
        row.registry_import = batch
        row.payment = payment
        row.distributor = owner
        row.save()
        batch.imported_count += 1
    batch.save(update_fields=['imported_count', 'duplicate_count'])
    return {'importId': batch.pk, 'imported': batch.imported_count, 'duplicates': batch.duplicate_count, 'skipped': skipped}


def reconciliation(row, payment, refunded=ZERO):
    if payment is None:
        return 'payment_not_found'
    if row.order_id_hint and row.order_id_hint != payment.order_id:
        return 'order_mismatch'
    if payment.currency != row.currency:
        return 'currency_mismatch'
    if payment.status != 'succeeded':
        return 'payment_not_succeeded'
    if row.kind == 'refund':
        return 'refund_exceeds_payment' if refunded > payment.amount else 'matched'
    if row.amount != payment.amount:
        return 'amount_mismatch'
    if row.fee is None or row.net is None:
        return 'fee_unknown'
    if row.amount - row.fee != row.net:
        return 'net_mismatch'
    return 'matched'


def build_report(start, end, distributor_id=None):
    """Payment-period sales and operation-period refunds, inclusive calendar bounds."""
    payments = Payment.objects.filter(provider='yookassa', currency='RUB', status='succeeded').select_related('order__client', 'order__distributor')
    entries = YooKassaRegistryOperation.objects.select_related('payment__order__client', 'payment__order__distributor')
    if distributor_id:
        payments = payments.filter(order__distributor_id=distributor_id)
        entries = entries.filter(Q(payment__order__distributor_id=distributor_id) | Q(payment__isnull=True, distributor_id=distributor_id))
    all_entries = list(entries)
    sales = {e.payment_id: e for e in all_entries if e.kind == 'payment' and e.payment_id}
    refunds = {}
    for entry in all_entries:
        if entry.kind == 'refund' and entry.payment_id:
            refunds[entry.payment_id] = refunds.get(entry.payment_id, ZERO) + entry.amount
    operations = []
    issues = []
    fees, refund_total, known_net, revenue = ZERO, ZERO, ZERO, ZERO
    unknown = 0
    orders = set()
    seen = set()

    def serialize(payment, date, amount, fee, refund, net, status, entry=None):
        order = payment.order if payment else None
        return dict(date=date.isoformat(), orderId=order.pk if order else (entry.order_id_hint if entry else None),
            client=order.client.company_name if order else None, distributor=order.distributor.name if order else None,
            distributorId=order.distributor_id if order else (entry.distributor_id if entry else None),
            paymentId=payment.pk if payment else None,
            providerPaymentId=payment.provider_payment_id if payment else entry.provider_payment_id,
            refundId=entry.provider_refund_id if entry else None,
            amount=str(amount), fee=str(fee) if fee is not None else None, refund=str(refund),
            net=str(net) if net is not None else None, status=status,
            registryAmount=str(entry.amount) if entry else None,
            kind=entry.kind if entry else 'payment')

    payments = payments.filter(Q(paid_at__gte=start, paid_at__lt=end) | Q(paid_at__isnull=True, created_at__gte=start, created_at__lt=end))
    for payment in payments:
        entry = sales.get(payment.pk)
        status = reconciliation(entry, payment) if entry else 'not_in_registry'
        verified = status == 'matched'
        fee, net = (entry.fee, entry.net) if verified else (None, None)
        revenue += payment.amount
        orders.add(payment.order_id)
        if verified:
            fees += fee
            known_net += net
        else:
            unknown += 1
        operation = serialize(payment, payment.paid_at or payment.created_at, payment.amount, fee, ZERO, net, status, entry)
        operations.append(operation)
        if status != 'matched':
            issues.append(operation)
        if entry:
            seen.add(entry.pk)
    for entry in all_entries:
        if not start <= entry.occurred_at < end or entry.pk in seen:
            continue
        payment = entry.payment
        status = reconciliation(entry, payment, refunds.get(entry.payment_id, ZERO))
        if entry.kind == 'payment' and payment and payment.status == 'succeeded':
            # Linked successful sales use the local paid_at, not the statement date.
            continue
        valid_refund = entry.kind == 'refund' and status == 'matched'
        if valid_refund:
            refund_total += entry.amount
            known_net -= entry.amount
        operation = serialize(payment, entry.occurred_at, ZERO if entry.kind == 'refund' else entry.amount,
            ZERO if valid_refund else None, entry.amount if entry.kind == 'refund' else ZERO,
            -entry.amount if valid_refund else None, status, entry)
        operations.append(operation)
        if status != 'matched':
            issues.append(operation)
    operations.sort(key=lambda row: row['date'], reverse=True)
    return dict(summary=dict(paidOrders=len(orders), sales=str(revenue), commission=str(fees) if not unknown else None,
        confirmedCommission=str(fees), refunds=str(refund_total), net=str(known_net) if not issues else None,
        confirmedNet=str(known_net), unreconciledPayments=unknown, discrepancies=len(issues),
        status='discrepancies' if issues else ('matched' if operations else 'empty')),
        operations=operations, discrepancies=issues)


def export_report(report, date_from, date_to):
    book = Workbook()
    summary = book.active
    summary.title = 'Итоги'
    summary.append(['Учет ЮKassa', f'{date_from} — {date_to}'])
    labels = {'paidOrders': 'Оплаченных заказов', 'sales': 'Продажи, RUB', 'commission': 'Комиссия с НДС, RUB',
              'refunds': 'Возвраты, RUB', 'net': 'К зачислению, RUB', 'confirmedCommission': 'Подтвержденная комиссия, RUB',
              'confirmedNet': 'Подтвержденная часть к зачислению, RUB', 'status': 'Статус сверки', 'discrepancies': 'Расхождений'}
    for key, label in labels.items():
        value = report['summary'][key]
        if key == 'status':
            value = STATUS_LABELS[value]
        summary.append([label, Decimal(value) if key in ('sales', 'commission', 'refunds', 'net', 'confirmedCommission', 'confirmedNet') and value is not None else value if value is not None else 'Не сверено'])
    summary.append(['Основа', 'Продажи: дата оплаты; возвраты: дата возврата. Только ЮKassa, RUB.'])
    headers = ['Дата', 'Заказ', 'Клиент', 'Дистрибьютор', 'ID платежа ЮKassa', 'Оплата, RUB', 'Комиссия с НДС, RUB', 'Возврат, RUB', 'К зачислению, RUB', 'Сверка', 'ID возврата', 'Сумма в реестре, RUB']
    keys = ['date', 'orderId', 'client', 'distributor', 'providerPaymentId', 'amount', 'fee', 'refund', 'net', 'status', 'refundId', 'registryAmount']
    for title, rows in [('Операции', report['operations']), ('Расхождения', report['discrepancies'])]:
        sheet = book.create_sheet(title)
        sheet.append(headers)
        for row in rows:
            sheet.append([Decimal(row[k]) if k in ('amount', 'fee', 'refund', 'net', 'registryAmount') and row[k] is not None else STATUS_LABELS[row[k]] if k == 'status' else row[k] for k in keys])
        sheet.auto_filter.ref = sheet.dimensions
        for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row) if rows else []:
            for index in (5, 6, 7, 8, 11):
                row[index].number_format = '#,##0.00;[Red]-#,##0.00'
    for sheet in book:
        sheet.freeze_panes = 'A2'
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='263238')
            cell.alignment = Alignment(wrap_text=True)
        sheet.row_dimensions[1].height = 32
        for column in sheet.columns:
            letter = get_column_letter(column[0].column)
            sheet.column_dimensions[letter].width = min(60, max(18, max(len(str(c.value or '')) for c in column) + 2))
        # Imported/user-controlled strings must stay literal, never Excel formulas.
        for row in sheet:
            for cell in row:
                if cell.data_type == 'f':
                    cell.data_type = 's'
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()
