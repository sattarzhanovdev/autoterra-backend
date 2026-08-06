"""Аналитика по покупкам: срезы и выгрузка.

Один источник расчётов для админки и для приложения менеджера. Если считать
в двух местах, цифры рано или поздно разойдутся, и доверия к отчёту не будет.

Покупка (``Purchase``) — это документ: клиент, дата, сумма, статус. Состав
документа — ``PurchaseItem``: SKU, количество, цена. Поэтому есть два уровня
среза: по документам (сколько закупили) и по позициям (что именно закупили).

В обороте участвуют только подтверждённые покупки: непроверенный документ ещё
не факт продажи, и складывать его в выручку нельзя.
"""

import csv
import io
import logging
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal

from django.db.models import Avg, Count, DecimalField, F, Q, Sum
from django.db.models.functions import Coalesce, TruncMonth
from django.http import HttpResponse
from django.utils import timezone

logger = logging.getLogger(__name__)

# Что считаем оборотом. Остальные статусы — заявки, а не продажи.
COUNTED_STATUSES = ("verified",)

MONEY = DecimalField(max_digits=14, decimal_places=2)


@dataclass
class Filters:
    """Набор фильтров отчёта. Пустое поле — «не ограничивать»."""

    date_from: date | None = None
    date_to: date | None = None
    distributor_id: str | None = None
    region_id: str | None = None
    sku: str | None = None
    status: str | None = None
    # Регионы, дальше которых менеджеру видеть нельзя. None — вся Россия.
    allowed_region_ids: list | None = field(default=None)

    @classmethod
    def from_request(cls, request, allowed_region_ids=None):
        return cls(
            date_from=_parse_date(request.GET.get("date_from")),
            date_to=_parse_date(request.GET.get("date_to")),
            distributor_id=_clean(request.GET.get("distributor")),
            region_id=_clean(request.GET.get("region")),
            sku=_clean(request.GET.get("sku")),
            status=_clean(request.GET.get("status")),
            allowed_region_ids=allowed_region_ids,
        )

    def as_dict(self):
        """Для подстановки обратно в форму и в ссылку выгрузки."""
        return {
            "date_from": self.date_from.isoformat() if self.date_from else "",
            "date_to": self.date_to.isoformat() if self.date_to else "",
            "distributor": self.distributor_id or "",
            "region": self.region_id or "",
            "sku": self.sku or "",
            "status": self.status or "",
        }


def _clean(value):
    value = (value or "").strip()
    return value or None


def _parse_date(value):
    value = (value or "").strip()
    if not value:
        return None
    for fmt in ("%Y-%m-%d", "%d.%m.%Y"):
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            continue
    logger.info("Не разобрана дата в фильтре отчёта: %r", value)
    return None


# ── Выборки ───────────────────────────────────────────────────────────────────

def purchases(filters: Filters):
    """Документы закупок с наложенными фильтрами."""
    from api.models import Purchase

    qs = Purchase.objects.select_related("client", "client__region", "distributor")

    # Статус: по умолчанию только подтверждённые — это и есть оборот.
    if filters.status:
        qs = qs.filter(status=filters.status)
    else:
        qs = qs.filter(status__in=COUNTED_STATUSES)

    if filters.date_from:
        qs = qs.filter(date__gte=filters.date_from)
    if filters.date_to:
        qs = qs.filter(date__lte=filters.date_to)
    if filters.distributor_id:
        qs = qs.filter(distributor_id=filters.distributor_id)
    # Регион — внешний ключ, поэтому фильтруем по id, а не по названию.
    if filters.region_id:
        qs = qs.filter(client__region_id=filters.region_id)
    if filters.allowed_region_ids is not None:
        qs = qs.filter(client__region_id__in=filters.allowed_region_ids)
    if filters.sku:
        # Документ попадает в отчёт, если в нём есть нужный SKU.
        qs = qs.filter(items__sku__iexact=filters.sku).distinct()

    return qs


def items(filters: Filters):
    """Позиции закупок — тот же фильтр, но на уровне товаров."""
    from api.models import PurchaseItem

    qs = PurchaseItem.objects.filter(purchase__in=purchases(filters).values("pk"))
    if filters.sku:
        # На уровне позиций нужен именно этот SKU, а не весь документ целиком.
        qs = qs.filter(sku__iexact=filters.sku)
    return qs.select_related("purchase", "purchase__client")


# ── Срезы ─────────────────────────────────────────────────────────────────────

def totals(filters: Filters) -> dict:
    """Итоги: сколько документов, на какую сумму, средний чек."""
    qs = purchases(filters)
    agg = qs.aggregate(
        count=Count("id"),
        amount=Coalesce(Sum("total_amount"), Decimal(0), output_field=MONEY),
        average=Coalesce(Avg("total_amount"), Decimal(0), output_field=MONEY),
    )
    agg["clients"] = qs.values("client_id").distinct().count()
    agg["quantity"] = items(filters).aggregate(
        total=Coalesce(Sum("quantity"), 0)
    )["total"]
    return agg


def by_sku(filters: Filters, limit=20):
    """Что закупают: топ позиций по сумме."""
    return list(
        items(filters)
        # Сумму позиции считаем до группировки: если сложить price*quantity
        # прямо в annotate, F("quantity") подхватит одноимённую аннотацию
        # вместо поля, и запрос упадёт.
        .annotate(line_total=F("price") * F("quantity"))
        .values("sku", "name")
        .annotate(
            quantity=Coalesce(Sum("quantity"), 0),
            amount=Coalesce(Sum("line_total", output_field=MONEY), Decimal(0), output_field=MONEY),
            documents=Count("purchase", distinct=True),
        )
        .order_by("-amount")[:limit]
    )


def by_client(filters: Filters, limit=20):
    """Кто закупает: топ клиентов по сумме."""
    return list(
        purchases(filters)
        .values("client__company_name", "client__inn", "client__region__name")
        .annotate(
            documents=Count("id"),
            amount=Coalesce(Sum("total_amount"), Decimal(0), output_field=MONEY),
        )
        .order_by("-amount")[:limit]
    )


def by_region(filters: Filters):
    """Срез по регионам — видно, где рынок живой, а где нет."""
    return list(
        purchases(filters)
        .values("client__region__name")
        .annotate(
            documents=Count("id"),
            clients=Count("client_id", distinct=True),
            amount=Coalesce(Sum("total_amount"), Decimal(0), output_field=MONEY),
        )
        .order_by("-amount")
    )


def by_distributor(filters: Filters):
    return list(
        purchases(filters)
        .values("distributor__name")
        .annotate(
            documents=Count("id"),
            amount=Coalesce(Sum("total_amount"), Decimal(0), output_field=MONEY),
        )
        .order_by("-amount")
    )


def by_month(filters: Filters):
    """Помесячная динамика — для графика."""
    return list(
        purchases(filters)
        .annotate(month=TruncMonth("date"))
        .values("month")
        .annotate(
            documents=Count("id"),
            amount=Coalesce(Sum("total_amount"), Decimal(0), output_field=MONEY),
        )
        .order_by("month")
    )


def report(filters: Filters) -> dict:
    """Полный отчёт одним вызовом — им пользуются и админка, и приложение."""
    return {
        "totals": totals(filters),
        "bySku": by_sku(filters),
        "byClient": by_client(filters),
        "byRegion": by_region(filters),
        "byDistributor": by_distributor(filters),
        "byMonth": by_month(filters),
    }


# ── Выгрузка ──────────────────────────────────────────────────────────────────

# Выгружаем позициями, а не документами: только так в отчёте виден SKU, ради
# которого аналитика и затевается.
COLUMNS = [
    ("Дата", lambda i: i.purchase.date.strftime("%d.%m.%Y") if i.purchase.date else ""),
    ("Документ", lambda i: i.purchase.document_number),
    ("Статус", lambda i: i.purchase.get_status_display()),
    ("Клиент", lambda i: i.purchase.client.company_name),
    ("ИНН", lambda i: i.purchase.client.inn),
    ("Регион", lambda i: i.purchase.client.region.name if i.purchase.client.region_id else ""),
    ("Дистрибьютор", lambda i: i.purchase.distributor.name if i.purchase.distributor_id else ""),
    ("Артикул", lambda i: i.sku),
    ("Товар", lambda i: i.name),
    ("Категория", lambda i: i.category or ""),
    ("Бренд", lambda i: i.brand or ""),
    ("Количество", lambda i: i.quantity),
    ("Цена", lambda i: float(i.price or 0)),
    ("Сумма позиции", lambda i: float((i.price or 0) * (i.quantity or 0))),
    ("Сумма документа", lambda i: float(i.purchase.total_amount or 0)),
]


def export_rows(filters: Filters, queryset=None):
    """Строки выгрузки. Отдельная функция — её же проверяют тесты.

    ``queryset`` позволяет выгрузить уже отобранные покупки — так работает
    действие в списке админки, где документы отмечают галочками.
    """
    rows = _items_for(filters, queryset)
    for item in rows.order_by("-purchase__date", "purchase__id"):
        yield [extract(item) for _, extract in COLUMNS]


def _items_for(filters: Filters, queryset=None):
    from api.models import PurchaseItem

    if queryset is None:
        return items(filters)
    qs = PurchaseItem.objects.filter(purchase__in=queryset.values("pk"))
    if filters.sku:
        qs = qs.filter(sku__iexact=filters.sku)
    return qs.select_related("purchase", "purchase__client")


def _filename(extension: str) -> str:
    stamp = timezone.localtime().strftime("%Y-%m-%d_%H-%M")
    return f"autoterra_purchases_{stamp}.{extension}"


def to_csv(filters: Filters, queryset=None) -> HttpResponse:
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";")
    writer.writerow([title for title, _ in COLUMNS])
    for row in export_rows(filters, queryset):
        writer.writerow(row)

    # BOM — иначе Excel открывает кириллицу кракозябрами.
    response = HttpResponse(
        "﻿" + buffer.getvalue(), content_type="text/csv; charset=utf-8"
    )
    response["Content-Disposition"] = f'attachment; filename="{_filename("csv")}"'
    return response


def to_xlsx(filters: Filters, queryset=None) -> HttpResponse:
    from api.services.exports import ExportUnavailable

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
    except ImportError as exc:  # pragma: no cover - зависит от окружения
        raise ExportUnavailable(
            "Для выгрузки в Excel нужен openpyxl: pip install openpyxl"
        ) from exc

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Покупки"

    for column, (title, _) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=1, column=column, value=title)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    for row_index, row in enumerate(export_rows(filters, queryset), start=2):
        for column, value in enumerate(row, start=1):
            sheet.cell(row=row_index, column=column, value=value)

    # Ширина по заголовку: без этого «Дистрибьютор» и «Сумма документа» режутся.
    for column, (title, _) in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=column).column_letter].width = (
            max(12, len(title) + 4)
        )
    sheet.freeze_panes = "A2"

    buffer = io.BytesIO()
    workbook.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{_filename("xlsx")}"'
    return response


FORMATS = {"csv": to_csv, "xlsx": to_xlsx}


def export(filters: Filters, fmt: str) -> HttpResponse:
    return export_queryset(None, filters, fmt)


def export_queryset(queryset, filters: Filters, fmt: str) -> HttpResponse:
    """Выгрузка либо по фильтрам, либо по уже отобранным документам."""
    exporter = FORMATS.get((fmt or "xlsx").lower())
    if exporter is None:
        from api.services.exports import ExportUnavailable

        raise ExportUnavailable(f"Неизвестный формат выгрузки: {fmt}")
    return exporter(filters, queryset)
